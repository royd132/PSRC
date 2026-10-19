import json

import pytest
from openpyxl import load_workbook

from formal.config import MODEL_NAMES
from formal.fm_table import TABLE_MODELS, build_fm_comparison_workbook


DATASETS = ("dkasc_site31", "dkasc_site1a", "pvod_station02", "hkust")
MODELS = MODEL_NAMES


def _result(dataset, model, rmse):
    return {
        "job": {"dataset": dataset, "horizon": 16, "model": model, "seed": 2026},
        "test": {
            "rmse_physical": rmse,
            "mae_physical": rmse / 2,
            "mbe_physical": rmse / 10,
            "r2": 1 - rmse / 10,
        },
        "parameter_source": "locked",
        "parameters": {},
        "epochs_requested": 100,
        "best_epoch": 10,
        "epochs_ran": 20,
        "early_stopped": True,
        "stop_reason": "patience",
        "patience": 10,
        "seconds": 1.0,
    }


def _fm_report(results):
    by_key = {(r["job"]["dataset"], r["job"]["model"]): r for r in results}
    stations = {}
    for dataset in DATASETS:
        psrc = by_key[(dataset, "psrc")]["test"]
        stations[dataset] = {
            "chosen_prior": "locked",
            "locked_source": "validation",
            "locked_hyperparameters": {
                "lr": 0.0003,
                "c": 1.0,
                "weight_decay": 0.0,
                "huber_beta_std": 1.0,
                "lambda_energy": 0.0,
                "lambda_smooth": 0.0,
            },
            "alpha_from_validation": 1.0,
            "A_psrc": dict(psrc),
            "FM_real": {
                "rmse_physical": psrc["rmse_physical"] - 0.1,
                "mae_physical": psrc["mae_physical"] - 0.05,
                "mbe_physical": 0.0,
                "r2": psrc["r2"] + 0.01,
            },
            "FM_shuffled": {"rmse_physical": psrc["rmse_physical"] - 0.05},
            "real_beats_psrc": True,
        }
    return {
        "per_station": stations,
        "protocol": "single test unblinding; adapters loaded from locked development artefacts; alpha* fixed from validation",
        "test_read_once": True,
    }


def test_fm_table_builds_locked_18_method_workbook(tmp_path):
    results = [
        _result(dataset, model, 1.0 + index / 10)
        for dataset in DATASETS
        for index, model in enumerate(MODELS)
    ]
    report = tmp_path / "fm.json"
    report.write_text(json.dumps(_fm_report(results)), encoding="utf-8")
    output = build_fm_comparison_workbook(results, report, tmp_path / "table.xlsx")

    workbook = load_workbook(output, data_only=True)
    comparison = workbook["Comparison"]
    # Ours + 17 public baselines (PSRC arm omitted from the presented table);
    # three metrics per model: RMSE / MAE / R2.
    assert comparison.max_column == 2 + 18 * 3
    assert comparison["C3"].value == "Ours"
    assert comparison["C4"].value == "RMSE"
    assert comparison["D4"].value == "MAE"
    assert comparison["E4"].value == "R2"
    assert "PSRC" not in [
        cell.value for row in comparison.iter_rows(min_row=3, max_row=4) for cell in row
    ]
    # PSRC synthetic RMSE is 1.0; Ours is 0.1 better at site 31.
    assert comparison["C5"].value == pytest.approx(0.9)
    assert comparison["C5"].font.bold is True
    assert comparison["C5"].font.color.rgb.endswith("C00000")
    # First public baseline (TimeMixer) starts at column F.
    assert comparison["F5"].value == pytest.approx(1.1)
    assert workbook["Run Metadata"].max_row == 1 + 4 * 18
    assert "Model Parameters" in workbook.sheetnames


def test_fm_table_rejects_psrc_mismatch(tmp_path):
    results = [
        _result(dataset, model, 1.0 + index / 10)
        for dataset in DATASETS
        for index, model in enumerate(MODELS)
    ]
    payload = _fm_report(results)
    payload["per_station"]["hkust"]["A_psrc"]["rmse_physical"] = 99.0
    report = tmp_path / "fm.json"
    report.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(ValueError, match="alignment failed"):
        build_fm_comparison_workbook(results, report, tmp_path / "table.xlsx")
